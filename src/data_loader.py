"""集中发现并读取赛题附件中的节点与运输无人机参数。"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict

from openpyxl import load_workbook


@dataclass(frozen=True)
class SourceFiles:
    problem_docx: Path
    nodes_xlsx: Path
    transport_xlsx: Path
    dem_tif: Path
    spatial_pdf: Path


@dataclass(frozen=True)
class Node:
    node_id: str
    name: str
    longitude_deg: float
    latitude_deg: float
    ground_elevation_m: float
    population: int | None = None


@dataclass(frozen=True)
class TransportModel:
    model_id: str
    name: str
    empty_mass_with_battery_kg: float
    max_payload_kg: float
    cargo_volume_m3: float
    cruise_speed_m_s: float
    empty_range_m: float
    full_range_m: float
    usable_energy_kwh: float
    return_reserve_fraction: float
    preparation_time_s: float
    loading_time_per_box_s: float
    handover_base_time_s: float
    handover_time_per_box_s: float
    climb_speed_m_s: float
    descent_speed_m_s: float
    climb_efficiency: float
    descent_efficiency: float


@dataclass(frozen=True)
class ProjectData:
    files: SourceFiles
    nodes: Dict[str, Node]
    models: Dict[str, TransportModel]


def _find_unique(root: Path, filename: str) -> Path:
    matches = sorted(path for path in root.rglob(filename) if path.is_file())
    if not matches:
        raise FileNotFoundError(f"在 {root} 下未找到附件：{filename}")
    if len(matches) > 1:
        raise RuntimeError(f"附件名称不唯一：{filename} -> {matches}")
    return matches[0]


def discover_source_files(problem_root: Path) -> SourceFiles:
    root = Path(problem_root).resolve()
    if not root.exists():
        raise FileNotFoundError(f"赛题根目录不存在：{root}")
    return SourceFiles(
        problem_docx=_find_unique(root, "山区洪涝灾害下无人机运输与通信协同优化.docx"),
        nodes_xlsx=_find_unique(root, "调度中心与服务区.xlsx"),
        transport_xlsx=_find_unique(root, "运输无人机数据.xlsx"),
        dem_tif=_find_unique(root, "镇龙乡及周边30米DEM.tif"),
        spatial_pdf=_find_unique(root, "镇龙乡地理空间数据说明.pdf"),
    )


def load_nodes(path: Path) -> Dict[str, Node]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["数据"]
    nodes: Dict[str, Node] = {}
    for row in worksheet.iter_rows(values_only=True):
        node_id = row[0]
        if not isinstance(node_id, str) or not (
            node_id == "O01" or node_id.startswith("S")
        ):
            continue
        nodes[node_id] = Node(
            node_id=node_id,
            name=str(row[1]),
            longitude_deg=float(row[2]),
            latitude_deg=float(row[3]),
            ground_elevation_m=float(row[4]),
            population=int(row[5]) if len(row) > 5 and row[5] is not None else None,
        )
    workbook.close()
    expected = {"O01", *(f"S{i:03d}" for i in range(1, 16))}
    missing = sorted(expected - set(nodes))
    if missing:
        raise ValueError(f"节点表缺少节点：{missing}")
    return nodes


def load_transport_models(path: Path) -> Dict[str, TransportModel]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["数据"]
    models: Dict[str, TransportModel] = {}
    for row in worksheet.iter_rows(min_row=3, max_row=5, values_only=True):
        if row[0] not in {"A", "B", "C"}:
            continue
        reserve = float(row[9])
        models[str(row[0])] = TransportModel(
            model_id=str(row[0]),
            name=str(row[1]),
            empty_mass_with_battery_kg=float(row[2]),
            max_payload_kg=float(row[3]),
            cargo_volume_m3=float(row[4]),
            cruise_speed_m_s=float(row[5]),
            empty_range_m=float(row[6]),
            full_range_m=float(row[7]),
            usable_energy_kwh=float(row[8]),
            return_reserve_fraction=reserve / 100.0 if reserve > 1 else reserve,
            preparation_time_s=float(row[10]),
            loading_time_per_box_s=float(row[11]),
            handover_base_time_s=float(row[12]),
            handover_time_per_box_s=float(row[13]),
            climb_speed_m_s=float(row[14]),
            descent_speed_m_s=float(row[15]),
            climb_efficiency=float(row[16]),
            descent_efficiency=float(row[17]),
        )
    workbook.close()
    if set(models) != {"A", "B", "C"}:
        raise ValueError(f"运输无人机参数不完整，实际读取：{sorted(models)}")
    return models


def load_project_data(problem_root: Path) -> ProjectData:
    files = discover_source_files(problem_root)
    return ProjectData(
        files=files,
        nodes=load_nodes(files.nodes_xlsx),
        models=load_transport_models(files.transport_xlsx),
    )
